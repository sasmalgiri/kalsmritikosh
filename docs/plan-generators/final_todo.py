import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

REPO = "/Users/shirshendusasmal/Downloads/towerlogicstrategicgame/original photo and ai assistant/Atlas chronica memora"
MD = REPO + "/docs/BEAT_DEVONTHINK_FINAL_TODO.md"
XLSX = "/Users/shirshendusasmal/Downloads/DEVONthink_vs_Kalsmritikosh_Gap_Matrix.xlsx"

DT, ML, BOTH, MOAT, GATE = "DEVONthink gap", "mailin port", "DT gap + mailin port", "Moat (DT can't)", "Ship gate"
MKT, MKTML = "Market gap", "Market gap + mailin port"
INNOV = "Innovation (no competitor has it)"

# (id, wave, title, source, why, how (incl. mailin file to port), done means, effort, priority, owner needed)
T = [
 # ---------------- Wave 0 — the toe-to-toe proof (both apps offline)
 ("W0-M", 0, "Module matrix: every feature behind its own on/off switch", BOTH,
  "Owner request 2026-10-15: make every new feature easy to turn on and off. Lets a misbehaving feature be switched off instead of reverted, and lets Private/Online gating apply per feature.",
  "Extend the existing KnowledgeModule switchboard (one enum, persisted switch, declared dependencies, AI gate, auto-listed in Settings ▸ Modules): one case per switchable plan item tagged with its plan ID, 'not implemented' until its code lands; a new requiresOnline gate (runs only in Online mode AND when the build has the network helper); MODULE_MATRIX.md generated from the enum by a test, like CAPABILITY_MATRIX.md. Invariants that must never be switchable (send preview, outbound log, evidence gate) are listed as such.",
  "Every switchable item in this plan has a module; the matrix doc regenerates and fails CI on drift; online modules read 'not available in this build' until O-2.", "S", "P0", "No"),
 ("W0-A", 0, "Head-to-head benchmark vs DEVONthink — both modes", MOAT,
  "Owner decision 2026-10-15: go toe to toe, with Private (offline) and Online (own API key) modes. Without a measured win, every claim is marketing.",
  "One corpus of 2,000-5,000 real-shaped files (PST/mbox mail, PDFs incl. scans, Office docs, chat exports); 50 questions: fact lookup, date/when, who-said-what, timeline, contradiction, and 10 whose answer is NOT in the corpus. Round 1 PRIVATE: Wi-Fi off; DEVONthink 4.4 Pro with its best local model (Ollama/LM Studio) vs Kalsmritikosh Private mode. Round 2 ONLINE: both apps use the SAME cloud model with the user's key (e.g. Claude). Score per answer: correct, citation opens the right passage, invented claims, honest refusal; plus setup time, ingest time, time per answer, peak RAM, and (online) bytes sent. Corpus, questions and scoring sheet live in the repo and rerun every release.",
  "Scorecard for both rounds published in docs/ with raw answers from both apps; rerun on every release candidate.", "M", "P0", "Yes — DEVONthink trial, local model, API key; run DEVONthink's side"),
 ("W0-B", 0, "Answer quality on real data at benchmark scale", MOAT,
  "Today's live ledger holds ~52 documents; the last real-data audit found passage-quality answers and a missed fact. The moat must hold at thousands of files.",
  "Ingest the W0-A corpus; fix every lost question at its root (extraction, retrieval, composer), never by special-casing the question. Track the scorecard after each fix.",
  "Kalsmritikosh wins W0-A on correct answers AND correct citations, with fewer invented claims; refusals on all 10 absent-answer questions.", "L", "P0", "No"),
 ("W0-C", 0, "Private mode is provably offline", BOTH,
  "Online mode must not weaken the privacy proof. The main app keeps network permission OFF; only a separate helper process has it (see O-2).",
  "Port mailin NO_NETWORK_PROOF.md approach: a script reads the signed Release app's entitlements and fails if the MAIN app has any network key; the helper's entitlements are checked to contain network-client only. A measured run in Private mode shows zero sockets and the helper never launched.",
  "Script runs in CI on the archived Release build; result linked from the App Store text and Settings.", "S", "P0", "No"),
 ("W0-1", 0, "Recorded scale run on owner hardware", GATE,
  "DEVONthink's biggest trust advantage is proven large databases; we may only advertise a measured figure.",
  "Run the P9.2 scale run on the new index selector; adopt mailin's ProcessFootprint + AdaptiveBatchController (ArchiveCore) to keep memory flat — today's out-of-memory warning shows the need.",
  "A tested figure recorded in SHIP_DECISIONS.md; memory stays flat across the run.", "L", "P0", "Yes — hardware run"),
 ("W0-2", 0, "Visual/binary redaction verified (F7)", GATE,
  "Release gate; redaction is a feature DEVONthink does not have — it must be airtight.",
  "Finish the F7 gate; compare against mailin RedactionEngine for PDF burn-in behaviour.",
  "Exported PDFs contain no recoverable redacted text or image data (tests + manual check).", "M", "P0", "No"),
 ("W0-3", 0, "Owner self-test + clean-machine install", GATE,
  "Ship gates in SHIP_DECISIONS §3.",
  "~20 real questions on the owner archive with openable citations; install on a Mac that never had Xcode.",
  "Checklist ticked in PRODUCTION_STATUS.md.", "M", "P0", "Yes"),

 # ---------------- Wave O — Online mode (own API key), Private stays the default
 ("O-1", 5, "Change control: lift 'no network provider in release'", GATE,
  "SHIP_DECISIONS §1 locks no network in release; CLAUDE.md and the CI guard enforce it. Owner decision 2026-10-15 changes it.",
  "Dated change-control entry in SHIP_DECISIONS: two modes, Private default, Online = own-key cloud AI only (no local server, sync or MCP). Update ci/guards/release-configuration.sh to allow network-client ONLY in the helper's entitlements. PrivacyGate stays and becomes mode-aware; it is never removed.",
  "Entry dated and approved; guard fails if the main app gains a network key.", "S", "P0", "Yes — approve the entry"),
 ("O-2", 5, "Network helper process (XPC service)", BOTH,
  "Keeps the kernel-enforced proof for the main app (W0-C) while allowing Online mode.",
  "New XPC service target containing only the cloud HTTP client (code from Routing/Providers). The main app sends it a prepared request; the helper returns the response. Private mode never launches it.",
  "Main app entitlements: no network. Helper: network-client only. Online calls work; Private mode leaves the helper unlaunched (tested).", "M", "P0", "Yes — create the target in Xcode"),
 ("O-3", 5, "Two modes: Private (default) and Online (opt-in)", BOTH,
  "Owner choice: start fully offline; Online is an explicit opt-in with a visible badge and a one-click way back.",
  "Mode switch in Settings and toolbar; badge always visible in Online mode; switching back is instant and stops the helper. Optional: pin a case to Private so it never goes online.",
  "New user starts Private; mode survives relaunch; a Private-pinned case refuses to send.", "M", "P0", "No"),
 ("O-4", 5, "Bring-your-own-key cloud providers: Claude, OpenAI, Gemini", BOTH,
  "Matches DEVONthink's main AI option.",
  "Unstub CloudProvider and use CloudEndpointRegistry; API key in the Keychain; model names live only in Routing/Providers (CLAUDE.md invariant). Capability resolution still goes through PrivacyGate.",
  "Each provider answers a fixture question; a bad key or no network is reported clearly.", "M", "P0", "No"),
 ("O-5", 5, "Send preview, consent and an outbound log", MOAT,
  "Our edge over DEVONthink online: the user sees exactly what leaves, and it is recorded.",
  "Before each send: the exact passages, destination and size; PII redaction option (reuse the redaction engine); 'ask every time' or 'this session'. Every send is written to an outbound log with a signed receipt (W1-3). Only the passages for that question are sent, never whole documents.",
  "Sent bytes equal the previewed passages (byte compare in tests); the log lists every send.", "M", "P0", "No"),
 ("O-6", 5, "Same evidence gate for cloud answers", MOAT,
  "A stronger model must not bypass the claim–evidence contract.",
  "Cloud output goes through the same claim validation, generative-call budgets and quality strip; cloud text is never primary evidence.",
  "A cloud answer citing a passage it was not given is downgraded or refused (test).", "S", "P0", "No"),
 ("O-7", 5, "No silent fallback between modes", BOTH,
  "Owner rule: no silent fallbacks.",
  "If Online fails (no network, quota, bad key), say so and offer 'answer privately instead'; never switch on your own.",
  "Network-off test shows the message and no automatic retry in the other mode.", "S", "P1", "No"),
 ("O-8", 5, "Privacy policy, App Store labels, review notes", GATE,
  "Sending personal data to a third-party AI requires clear disclosure and permission under App Review rules.",
  "Update the privacy policy and App Store privacy answers (data goes to the provider the user chooses, with their key; we collect nothing); consent screen before first send; App Review notes explain both modes.",
  "Documents updated; consent shown once before the first send.", "S", "P0", "Yes — review legal text"),

 # ---------------- Wave 1 — quick wins, mostly ports (table stakes vs DEVONthink)
 ("W1-1", 1, "Advanced search syntax", BOTH,
  "DEVONthink power users live on Boolean/NEAR/field operators; we only expose plain FTS5.",
  "Port mailin ArchiveQueryCompiler (AND/OR/NOT, phrase, NEAR/n, wildcard, regex, field ops) to compile into our FTS5 + metadata filters (from:, date:, type:, entity:, tag:).",
  "Each operator has a failing-first test; unknown operators are reported, never silently treated as words (mailin found that exact bug).", "M", "P1", "No"),
 ("W1-2", 1, "Search coverage badge — 'where it matched, how much was indexed'", BOTH,
  "Honesty about coverage is our moat; DEVONthink never says what it has not indexed.",
  "Port mailin SearchCoverage + SearchMatchFieldChips; feed from our per-lane readiness (chunks/vectors/OCR).",
  "Every result list shows matched fields and indexed-of-total; a deferred source is named.", "S", "P1", "No"),
 ("W1-3", 1, "Signed import + export receipts", BOTH,
  "Our competitive wedge is verifiable receipts; mailin already signs them.",
  "Port ImportReceipt, ExportReceipt, ExportSigner (Ed25519, CryptoKit) and ExportFolderManifest; link into the existing ReceiptVerifierView.",
  "Every ingest and export ends in a receipt with hashes; a changed byte fails verification.", "M", "P1", "No"),
 ("W1-4", 1, "User tags + tag-scoped search", BOTH,
  "Tags are DEVONthink's core organiser; we have none.",
  "New versioned migration (tags, item_tags); port TagDisplayPolicy/TagOverridePersistence rules; tags usable as a search and Ask scope. Tags are user metadata — they never change evidence status.",
  "Tag, untag, rename, merge; tag: operator works; migration test pinned.", "M", "P1", "No"),
 ("W1-5", 1, "Deep links: kalsmritikosh://", BOTH,
  "DEVONthink item links are used in every notes app; mailin plans the same (E2b).",
  "URL scheme for document, exact locator (page/line), citation and saved answer; 'Copy link' everywhere; missing item says so. Info.plist via Xcode's tool.",
  "Links open the exact passage after relaunch; a link to a deleted/absent item reports it.", "S", "P1", "Yes — Info.plist/URL type in Xcode"),
 ("W1-6", 1, "Shortcuts / App Intents", BOTH,
  "DEVONthink has scripting + Shortcuts; we have no automation surface.",
  "Port the AppIntentShortcuts pattern: Ask (returns answer + citations), Search, Ingest folder, Open source, Export dossier; AppEntity for Document, Entity, Event, Tag.",
  "Each action runs from the Shortcuts app and `shortcuts run`; counts equal SQL.", "M", "P1", "No"),
 ("W1-7", 1, "Command palette + keyboard-first navigation", ML,
  "Power-user speed; mailin has it.",
  "Port CommandPaletteView + KeyboardShortcutOverlayView; register every screen and main action.",
  "Every main action reachable from ⌘⇧P; shortcut overlay lists them.", "S", "P2", "No"),
 ("W1-8", 1, "Move the ledger to an external disk", ML,
  "Large archives fill a 256 GB Mac (seen today); DEVONthink databases can live anywhere.",
  "Port ArchiveRelocator (verified copy, keep working while the disk is away) + StoragePlanner for space estimates.",
  "Ledger moved and verified by hash; app reports clearly when the disk is absent.", "M", "P1", "No"),
 ("W1-9", 1, "Import queue control", ML,
  "Long ingests need pause/resume/reorder; mailin has a full queue.",
  "Port ImportQueueView + BulkImportCoordinator behaviour onto our IngestCoordinator and source lease.",
  "Pause, resume, stop one source, reorder; resume never duplicates (reuse F04 lease tests).", "M", "P2", "No"),

 # ---------------- Wave I — innovation: things no competitor in the scan does
 ("I-1", 7, "Proof bundle: an answer anyone can verify without the app", INNOV,
  "Everyone shows citations; nobody lets a court, editor or client check them independently. Today our receipts use an unkeyed hash chain (proves consistency, not origin).",
  "Export any answer as one signed bundle: the answer, every quoted span with its exact locator, the SHA-256 of each source file, the ledger revision, and an Ed25519 signature (key in the Keychain; builds on W1-3 and ReceiptVerifierView). Ship a tiny open verifier (a single HTML file that runs offline in any browser) that re-checks quotes against the user's copies of the sources.",
  "Changing one character of a quote or source makes the verifier fail at that exact span; a valid bundle verifies on a Mac without Kalsmritikosh.", "M", "P1", "No"),
 ("I-2", 7, "'What is missing' map", INNOV,
  "Other tools only show what they found. Investigators win on what is absent. GapDetector already finds numbered runs, explicit references and reply headers that imply a missing item.",
  "Surface GapDetector results as a map per case/question: missing numbers in a series (invoices, exhibits, statements by month), replies to emails we do not have (N-2), quiet periods with no documents, referenced attachments not present. Each gap shows the pattern that implies it and stays low-confidence (absence is inferred, never a fact).",
  "Fixture with a missing invoice number, a missing month and an orphan reply shows exactly those three gaps with reasons; no gap is invented from a sparse sequence.", "M", "P1", "No"),
 ("I-3", 7, "Corroboration meter on every claim", INNOV,
  "'Cited' is not 'proven': five copies of one email are one source. Source-independence counting exists in retrieval; users never see it.",
  "Next to each claim: number of INDEPENDENT sources (copies, forwards and near-duplicates from W2-4 count once), their kinds, and 'single source — treat as a lead' when there is only one (the FactCheckMemo discipline, applied everywhere).",
  "A claim backed by an email and its forward shows 1 independent source; adding an unrelated document makes it 2.", "S", "P1", "No"),
 ("I-4", 7, "Devil's-advocate check before an answer ships", INNOV,
  "No competitor actively looks for evidence AGAINST its own answer.",
  "After composing, run a disconfirmation search (negations, conflicting dates/amounts, opposing parties; reuse ACHAnalysis) and show 'searched N sources for contradicting evidence: none found' or the conflict with both sources. Counts against the existing generative-call budget; deterministic where possible.",
  "Planting a contradicting document turns a confident answer into a surfaced conflict; with none planted the 'searched N sources' line appears.", "M", "P1", "No"),
 ("I-5", 7, "Time travel: 'what did we know on date X?'", INNOV,
  "Journalists and investigators need to know what was knowable when. The ledger carries document dates and ingest versions.",
  "An 'as of' control on Ask, Timeline and Dossier: answer using only evidence dated (or received) on or before X; a slider shows how the answer changed over time and which document changed it.",
  "On a fixture, the answer as of an earlier date omits facts first stated in later documents; the change point names the document.", "M", "P2", "No"),
 ("I-6", 7, "Fact-check my writing, live", INNOV,
  "Writers draft outside the evidence. FactCheckMemo does this for journalists after the fact; nobody does it while you type.",
  "In the notebook/report editor, each sentence is checked against the ledger: supported (with source), contradicted (shows the conflicting source), or unsupported. Works offline; only the user's own text and ledger are used.",
  "Typing a wrong date from a fixture flags it red with the source; a correct one turns green.", "M", "P1", "No"),
 ("I-7", 7, "Measured accuracy on YOUR archive + calibrated confidence", INNOV,
  "Every AI tool says 'may be inaccurate'. None tells you how accurate it is on your own data.",
  "Extend AnswerHarness: generate checkable questions from deterministic ledger facts (dates, amounts, parties) plus questions that must be refused, answer them blind, and show a personal scorecard. Feed results into CalibrationStore so displayed confidence matches measured accuracy.",
  "On the benchmark corpus, the in-app score is within a few points of the W0-A manual score; '80% confident' answers are right about 80% of the time.", "M", "P1", "No"),
 ("I-8", 7, "Least-data online answers + privacy budget", INNOV,
  "Elephas masks personal data; nobody minimises WHAT is sent.",
  "In Online mode, answer everything possible locally first; send the cloud model only the minimal masked sub-question and passages it truly needs (with N-10 masking); show bytes and fields sent per answer and a running per-case privacy budget the user can cap or set to zero.",
  "For fixture questions answerable locally, nothing is sent; for the rest, sent text is smaller than the full-passage approach and matches the preview byte-for-byte.", "M", "P1", "No"),
 ("I-9", 7, "Statement drift: how a person's account changed", INNOV,
  "Contradiction detection compares documents; lawyers need to compare one PERSON over time.",
  "For a person, line up everything they said or wrote about a topic in date order (emails, transcripts, testimony) and flag shifts in dates, amounts, who-was-there and denials, each side cited.",
  "A fixture witness whose date changes between an email and a transcript is flagged with both quotes.", "M", "P2", "No"),
 ("I-10", 7, "Signed case slices: share without a server", INNOV,
  "Collaboration usually means a cloud. We can share a portable, verifiable slice instead.",
  "Export a redacted, signed, read-only subset of the ledger (chosen facts, events, entities, quoted spans) as one file another Kalsmritikosh — or the iPad viewer (W4-6) — opens; tampering is detected. No sync service, works in Private mode.",
  "A slice opens on a second Mac; editing one byte is detected; redacted text is absent from the file.", "M", "P2", "No"),
 ("I-11", 7, "Explain-the-path view for every answer", INNOV,
  "Users trust what they can follow. We already record the retrieval priority path (memory → timeline → entity → FTS → … → vector).",
  "A 'how this answer was found' panel: which lanes ran, what each found, what the evidence gate kept or dropped and why, with the generative calls used. Readable language, no internals.",
  "Every answer shows its path; a refused answer explains which lane found nothing.", "S", "P2", "No"),
 ("I-12", 7, "Living dossiers that update themselves", INNOV,
  "DEVONthink and Fenn store files; nobody keeps a profile of each person or company current as new files arrive.",
  "Per entity: key facts, timeline, relationships, contradictions, gaps and watchlist hits, rebuilt after each ingest with a 'what changed since you last looked' strip (reuses N-11 change detection and the existing DossierView).",
  "Ingesting a new fixture document updates the dossier and lists exactly the new facts.", "M", "P2", "No"),

 # ---------------- Wave 2 — evidence depth (where we pull far ahead)
 ("W2-1", 2, "Annotations and highlights as an overlay", DT,
  "DEVONthink annotates PDFs; we can't annotate anything.",
  "Overlay records (never edit the source) anchored to our exact locators; a note becomes HUMAN_CONFIRMED/CORRECTED evidence in the existing vocabulary; port mailin examiner-annotation model.",
  "Highlights survive re-ingest of an unchanged source; exports can include or omit them.", "L", "P1", "No"),
 ("W2-2", 2, "Email authenticity signals: DKIM, ARC, S/MIME, hidden timestamps", ML,
  "Neither DEVONthink nor most DMS tools verify email; mailin passes 166/168 of the ValiMail ARC suite.",
  "Port ForensicSignals (Gmail boundary/X-GM IDs, UUIDv1), DKIMVerifier/ARC chain, SMIMEHandler into Ingestion; store results as DETERMINISTICALLY_DERIVED facts shown in AuthenticityView. No DNS lookups (keys from the message or user-supplied).",
  "Metaspike worked examples decode; ARC 1- and 3-hop chains pass, tampered ones fail at the right hop.", "M", "P1", "No"),
 ("W2-3", 2, "PST reader: adopt mailin's MS-PST reader validated on 15 real PSTs", ML,
  "Our PST row is measured only at the field-mapping layer; mailin's reader matches readpst on 15 real Outlook files.",
  "Port PSTSpecReader behind a switch; run both readers on the same files and compare counts; keep ours as fallback.",
  "Item/attachment counts equal readpst on the public PST set; CAPABILITY_MATRIX evidence updated.", "M", "P1", "No"),
 ("W2-4", 2, "Near-duplicate detection across documents and emails", BOTH,
  "DEVONthink shows duplicates; near-duplicates (drafts, forwards) matter more for evidence.",
  "Port NearDuplicatesEngine (shingling/similarity); link near-dups as 'versions of' — never merge or delete (no-delete rule).",
  "Fixture with drafts and forwards groups correctly; independence counting treats near-dups as one source.", "M", "P2", "No"),
 ("W2-5", 2, "Legal production pack: Bates PDF, load files, legal holds", ML,
  "Lawyer persona; DEVONthink has stamps and Bates numbering, but no load files or legal holds.",
  "Port BatesEngine + BatesPDFRenderer, PreservationExporter (Concordance/Relativity DAT/OPT), legal hold model; integrate with PrivilegeLogStudioView and redaction.",
  "A production set opens in a review tool (owner check); Bates ranges continuous; holds block deletion.", "L", "P2", "Partly — review-tool check"),
 ("W2-6", 2, "Export breadth", ML,
  "mailin exports mbox, EML, MSG, PDF, TIFF, CSV, Word, Markdown, JSON, HTML with pre-flight and resume.",
  "Port ArchiveExportService renderers + ExportPreflight + ExportRunCenter for dossiers, timelines, entity lists and source sets.",
  "Every format round-trips its fixture; receipts attached (W1-3).", "M", "P2", "No"),
 ("W2-7", 2, "App lock (Touch ID) + at-rest encryption decision", DT,
  "DEVONthink offers encrypted databases; we rely on FileVault.",
  "Port BiometricLockManager now. Encryption: (a) AES-GCM for vault blobs with Keychain key — no new dependency; (b) whole-SQLite encryption needs SQLCipher, a third-party dependency — owner decision.",
  "App locks on sleep/idle; vault files unreadable without the key; decision recorded in SHIP_DECISIONS.", "M", "P1", "Yes — SQLCipher decision"),
 ("W2-8", 2, "Large single items by byte range", ML,
  "mailin keeps >100 MB messages by byte range instead of failing; avoids loading whole files.",
  "Compare mailin OffsetImportEngine/RawMessageAccess with our F01 streaming; adopt byte-range access for oversized items.",
  "A 200 MB message/attachment ingests with flat memory and an exact locator.", "M", "P2", "No"),

 # ---------------- Wave 3 — beat DEVONthink by a mile
 ("W3-1", 3, "Import a DEVONthink database (the switcher feature)", MOAT,
  "Lets DEVONthink users point us at their database and get a ledger in one step — DEVONthink cannot do the reverse.",
  "Read the database package's stored files in place (read-only, custody kept) and indexed folders; keep group path as a collection and Finder tags as tags. Verify the package layout on a real DEVONthink database before building.",
  "A sample DEVONthink database ingests with every file accounted for; groups and tags preserved; nothing written into the package.", "L", "P1", "Yes — a test DEVONthink database"),
 ("W3-2", 3, "Work alongside DEVONthink (export + links)", BOTH,
  "Be the evidence engine DEVONthink users add, as mailin plans (E2).",
  "'Send to DEVONthink': dossier/answer as Markdown or PDF with kalsmritikosh:// links via NSWorkspace.open (sandbox-safe; Finder fallback).",
  "Exported items open in DEVONthink; their links jump back to the exact passage.", "S", "P2", "No"),
 ("W3-3", 3, "Timeline and entity-graph visualisation", MOAT,
  "Our ledger has dated events and relationships; DEVONthink has nothing comparable. No graph UI exists yet.",
  "Interactive timeline (zoom, filter by entity/tag) and entity graph (Canvas, force layout) over InMemoryBondGraph; every node opens its evidence. Port contact-network ideas from ArchiveAnalyticsService.",
  "Graph of 1,000+ entities stays responsive; every edge cites its sources.", "L", "P1", "No"),
 ("W3-4", 3, "Proactive alerts: new contradictions, gaps and watched entities", MOAT,
  "DEVONthink is passive; we can tell users when a new file changes the story.",
  "Hook ingest completion → contradiction/gap detectors → alert list; port mailin keyword monitoring for watch terms.",
  "Ingesting a contradicting document raises an alert showing both sources.", "M", "P1", "No"),
 ("W3-5", 3, "Live smart collections + user rules", BOTH,
  "Smart groups and smart rules are DEVONthink staples.",
  "Saved queries become live collections with counts; port AutomationRulesEngine/recipes ('when a source from X arrives → tag/scope/alert'). JavaScript scripts (mailin ScriptEngine) only later, with its hostile-script guards.",
  "Rules run on ingest with a dry run; collections update without reload.", "M", "P2", "No"),
 ("W3-6", 3, "Capture: Share extension, Services menu, Safari web clipper", DT,
  "DEVONthink captures from anywhere; we only take files.",
  "New extension targets that drop items into the intake queue; the clipper saves page HTML/PDF locally — no network from the app.",
  "Share from Safari/Finder/Mail lands in the queue with custody recorded.", "M", "P2", "Yes — new targets in Xcode"),
 ("W3-7", 3, "Spotlight (opt-in)", BOTH,
  "DEVONthink items appear in Spotlight.",
  "Port SpotlightIndexer; OFF by default because it copies text into the system index — say so in Settings.",
  "Toggle on → items findable; toggle off → index removed.", "S", "P3", "No"),

 # ---------------- Wave N — new gaps from the market scan (2026-10-15)
 ("N-1", 6, "Fact chronology with issue tags + legal drafting", MKT,
  "Casefleet, CaseMap+ AI and Everlaw Storybuilder are built around facts linked to evidence and tagged by issue; we have no issue tags and no 'facts by issue' report.",
  "Issue tags on ledger facts/events (new migration); a 'Facts by issue' report; draft a statement of facts or deposition outline from selected facts, every sentence cited. Drafting uses the same evidence gate.",
  "Tag facts with issues, filter the chronology by issue, export a cited statement of facts.", "M", "P1", "No"),
 ("N-2", 6, "Communication analysis: who talks to whom + thread completeness", MKT,
  "Intella and Nuix show a communication graph, volume timeline and histogram, and flag a reply whose original email is missing from the case.",
  "Who-to-whom graph with message counts, volume over time, and a 'missing email in this thread' marker (In-Reply-To/References point to a message we do not have). Extends W3-3.",
  "On a fixture mailbox, every reply to an absent message is flagged; graph counts equal SQL.", "M", "P1", "No"),
 ("N-3", 6, "Watchlist cross-reference", MKT,
  "Aleph's 'killer feature': match your own list (people, companies, sanctions) against everything, including name variants.",
  "Import a CSV/JSON list; match against canonical entities and aliases with name-variant matching; each hit links to its evidence. Offline; the user supplies the list.",
  "A planted watchlist name in a different spelling is found with its source.", "M", "P1", "No"),
 ("N-4", 6, "Review coding with suggestions and reasons", MKTML,
  "Everlaw coding suggestions explain their reasoning; Intella and mailin have predictive coding.",
  "Code documents responsive / privileged / hot; suggestions learn from the reviewer's decisions and show why (port mailin PredictiveCodingEngine). Suggestions never change evidence status.",
  "After 50 coded documents, suggestions beat chance on a held-out fixture; each suggestion shows its reasons.", "M", "P2", "No"),
 ("N-5", 6, "Studio outputs from the ledger: briefing, FAQ, chronology, mind map, audio brief", MKT,
  "NotebookLM's Studio (briefings, FAQs, timelines, mind maps, audio overviews) is what non-experts love.",
  "One-click outputs built from the ledger, each line cited; mind map from topics and entities; an offline audio briefing read by on-device speech (no cloud voices).",
  "Each output regenerates from the ledger and every statement opens its source.", "M", "P2", "No"),
 ("N-6", 6, "Ask from anywhere: menu bar + global hotkey", MKTML,
  "Elephas Super Command and Fenn work from any app; mailin has a menu-bar search.",
  "Menu-bar search/ask panel and a global shortcut; answers open in the main window with citations (port mailin MenuBarSearchView).",
  "From any app, the hotkey opens the panel; an answer opens its source.", "S", "P2", "No"),
 ("N-7", 6, "Image intelligence: similar images and face grouping", MKT,
  "Fenn finds similar images and faces; investigators need 'where else does this photo appear'.",
  "On-device image feature prints for similarity; face grouping only after the user turns it on, with user-named groups and no automatic identity claims.",
  "A resized/cropped copy is found as similar; faces are grouped only with the switch on.", "M", "P2", "Owner — face feature policy"),
 ("N-8", 6, "Locations: photo GPS and addresses on a map", MKT,
  "Intella and Nuix have geolocation views.",
  "Read photo GPS and extracted addresses; Private mode shows a list with coordinates; the map view needs map tiles from the internet, so it is an Online-mode module.",
  "Fixture photos with GPS appear at the right places; Private mode makes no tile request.", "S", "P3", "No"),
 ("N-9", 6, "Import notes apps: Apple Notes, Obsidian, Notion exports", MKT,
  "Elephas indexes Apple Notes, Notion and Obsidian; these are where many users keep their thinking.",
  "Parsers for an Obsidian vault (wiki-links become relations), Notion export (Markdown/CSV) and Apple Notes (user-exported); each measured in CAPABILITY_MATRIX.",
  "Each format has a fixture and a FULL or honest PARTIAL row.", "M", "P2", "No"),
 ("N-10", 6, "Typed PII redaction before anything is sent online", MOAT,
  "Elephas redacts 28 entity types on-device before cloud calls; ours must match or beat it.",
  "Before an Online send, detect and mask names, emails, phones, addresses, IDs, account numbers using our own entities and redaction engine; the preview (O-5) shows masked text; answers are un-masked locally.",
  "A fixture with each PII type leaves the Mac masked (byte check); the answer reads normally.", "M", "P1", "No"),
 ("N-11", 6, "Question monitors: tell me when the answer changes", MOAT,
  "AnythingLLM runs scheduled agent jobs; Tana has live queries. Nobody tells you when new evidence changes an answer.",
  "Pin a question; after each ingest it re-runs and alerts when the answer, its citations or a contradiction changes, showing old vs new.",
  "Ingesting a file that changes a pinned answer raises one alert with both versions.", "M", "P2", "No"),
 ("N-12", 6, "Custom entity types and fields", MKT,
  "Tana supertags and Capacities typed objects let users define 'Contract: parties, value, start date'.",
  "User-defined types with fields; open-field extraction fills them where the value is written in a source; values show their evidence.",
  "Define a type, see it filled from fixtures, edit a value as HUMAN_CORRECTED.", "M", "P2", "No"),
 ("N-13", 6, "Explainable risk / priority score", MKT,
  "Aleph Pro and Nuix score risk to decide what to read first.",
  "Deterministic score per document/entity from contradictions, watchlist hits, authenticity failures, recency; every point of the score is explained.",
  "Score is reproducible and each component links to its evidence.", "S", "P3", "No"),

 # ---------------- Wave 4 — breadth and reach
 ("W4-1", 4, "OCR quality: promote images and scanned PDFs from PARTIAL", BOTH,
  "DEVONthink Pro OCR is strong; ours is measured PARTIAL.",
  "Port mailin E1 approach (Vision RecognizeTextRequest, one page in memory, OCR version column for a single re-pass); measure on real scans.",
  "CAPABILITY_MATRIX rows move to FULL on measured fixtures; searchable-PDF export.", "M", "P1", "No"),
 ("W4-2", 4, "Parser gaps: Keynote, ppt, iMessage, Safari/Chrome history, 7z/rar", DT,
  "Rows still PRESERVED-ONLY / NOT EXPANDED in the measured matrix.",
  "Structural parsers per format, each with a fixture so the matrix promotes it; 7z/rar only with a licence-clean decoder.",
  "Each row FULL or honestly PARTIAL in CAPABILITY_MATRIX.", "L", "P2", "Licence check for 7z/rar"),
 ("W4-3", 4, "Deleted-item recovery for mail stores", ML,
  "Investigator persona; mailin carves deleted PST/NSF/mbox items.",
  "Port RecoveryDetectors, MIMECarver, NSFDeletionStubs; recovered items carry a 'recovered' provenance and PARTIAL status, never mixed with live items.",
  "Emptied Deleted Items carved back on the public PST set; no invented items on clean files.", "L", "P2", "No"),
 ("W4-4", 4, "Review packs + team cases over a shared folder", ML,
  "Multi-reviewer work without a server; mailin signs every log entry.",
  "Port ReviewPack + TeamCase (Ed25519 signed logs, leases, Verify audit) onto our cases.",
  "Two reviewers on one shared folder see the same state; tampering reported at the exact entry.", "L", "P3", "Yes — second Mac / SMB"),
 ("W4-5", 4, "Languages: localized UI + multilingual ledger", BOTH,
  "DEVONthink is multilingual; we are English-only. mailin already ships 11 UI languages.",
  "String Catalog for the UI, reusing mailin translations for shared terms; v2 bge-m3 embedder per the multilingual plan.",
  "UI complete in the chosen languages; retrieval eval per language recorded.", "L", "P2", "Yes — language list"),
 ("W4-6", 4, "iPhone/iPad companion (read-only)", BOTH,
  "DEVONthink To Go; mailin is universal.",
  "Read-only ledger viewer from an exported, signed snapshot — no sync service, consistent with no-network.",
  "Open a snapshot on iPad; citations open their source.", "L", "P3", "Yes — scope"),
 ("W4-7", 4, "Accessibility audit + full UI crawl", ML,
  "mailin's audit process found real defects; App Store expectation.",
  "Reuse mailin's AccessibilityAuditUITests approach over Ask, Sources, Timeline and Settings.",
  "Audit green on every main surface.", "M", "P2", "Yes — presence for audit"),
 ("W4-8", 4, "User manual + 'Coming from DEVONthink' guide", DT,
  "DEVONthink has an extensive manual and community.",
  "Manual per persona with recipes; a switcher page mapping DEVONthink concepts (groups, tags, smart groups, See Also) to ours.",
  "Published with the release; linked from the in-app guide.", "M", "P2", "Owner review"),
 ("W4-9", 4, "Pricing tiers decision", ML,
  "mailin has Free/Personal/Professional via StoreKit; Kalsmritikosh plans one-time $29–49.",
  "Market scan: DEVONthink and Fenn sell a licence with one year of updates; Elephas ($6.99-29.99/month), Casefleet (~$30/user/month) and NotebookLM sell subscriptions. Options: one-time + paid yearly updates, or Free/Personal/Pro via mailin's TierAccess pattern.",
  "Decision recorded in SHIP_DECISIONS change log.", "S", "P2", "Yes — pricing"),
]


COMPETITORS = [
 ("DEVONthink 4.4", "Mac document manager", "$99 / $199 / $499, 1 yr updates", "Cloud or local AI chat, MCP server, versioning, revision-proof DBs, Bates, smart rules, sync, iOS", "No evidence gate, timelines, entities or contradiction checks", "Main rival; Waves 1-4", "https://www.devontechnologies.com/apps/devonthink"),
 ("Fenn", "Private AI file search (Mac)", "Lifetime + 1 yr updates", "On-device index of 50+ types, agentic search, jump to exact page/timestamp, faces and similar images, transcription, auto-organise, iPhone app", "Search and chat, no ledger, no evidence checking", "Closest consumer rival; N-6, N-7", "https://www.usefenn.com/"),
 ("Elephas", "AI knowledge assistant (Mac/iOS)", "$6.99-29.99 / month", "Offline toggle, own-key cloud models, PII redaction (28 types) before cloud, system-wide hotkey, Notes/Notion/Obsidian indexing", "Generic RAG; citations only", "N-6, N-9, N-10", "https://elephas.app/"),
 ("NotebookLM", "Cloud notebook (Google)", "Free / Plus / Pro tiers", "Cited chat, audio/video overviews, mind maps, briefing docs, FAQs, timelines", "Cloud only; 50-600 sources per notebook; usage quotas", "N-5", "https://fast.io/resources/notebooklm-limits/"),
 ("AnythingLLM", "Open-source local RAG desktop", "Free (MIT) + Pro", "Local models, workspaces, agents, MCP, scheduled jobs", "Answer quality depends on the chosen model; no ledger", "N-11", "https://docs.anythingllm.com/installation-desktop/overview"),
 ("Obsidian + AI plugins", "Notes app + Smart Connections / Copilot", "Free + plugin subscriptions", "Semantic links, vault chat, local or cloud models", "Only sees the vault; no AI built in", "N-9", "https://www.shadow.do/blog/best-ai-plugins-for-obsidian-2026"),
 ("Tana / Capacities / Heptabase", "Structured notes", "~$10-14 / month", "Typed objects and fields, live queries, MCP, canvas", "Notes you write, not archives you ingest", "N-12", "https://www.dexi.net/compare/tana-vs-capacities"),
 ("Intella (Vound)", "Desktop eDiscovery / investigation", "Per-licence (enterprise)", "Email threading with missing-email detection, social graph, timeline, geolocation, near-dups, predictive coding", "Expensive; review tool, not answers", "N-2, N-4, N-8", "https://www.vound-software.com/ediscovery"),
 ("Nuix Neo", "Enterprise investigations", "Enterprise", "1,000+ file types, entity extraction at ingest, link analysis, risk scoring, full audit trail", "Enterprise price and scale", "N-13", "https://www.nuix.com/solutions/fraud-investigations"),
 ("Everlaw", "Cloud eDiscovery + AI", "Enterprise", "Deep Dive cited Q&A over millions of docs, coding suggestions with reasons, Storybuilder timelines and depositions; Claude/Copilot integrations", "Cloud only; legal teams", "N-1, N-4", "https://www.everlaw.com/product/"),
 ("Casefleet / CaseMap+ AI", "Litigation fact chronology", "~$30-40 / user / month (Casefleet)", "Facts linked to evidence, issue tags, 'facts by issue' reports, AI proposes facts, drafts motions", "Cloud; you build the chronology", "N-1", "https://www.casefleet.com/comparisons/casemap-alternative"),
 ("Aleph Pro (OCCRP)", "Investigative journalism platform", "Free for journalists (access controlled)", "1B+ public records, cross-referencing with name variants, watchlists, knowledge graph", "Hosted; your files must be uploaded", "N-3", "https://gijn.org/resource/using-aleph/"),
 ("Rewind / Limitless", "Screen memory (Mac)", "Shut down (Meta acquisition, Dec 2025)", "—", "Gone — a warning about cloud-dependent memory apps", "Trust: we are local and our ledger is yours", "https://luci.memories.ai/blog/rewind-ai-shut-down-on-device-replacement"),
]

WAVES = {0: "Wave 0 — the toe-to-toe proof + v1 ship gates",
         1: "Wave 1 — quick wins, mostly ports from mailin (DEVONthink table stakes)",
         2: "Wave 2 — evidence depth (pull far ahead)",
         3: "Wave 3 — beat DEVONthink by a mile",
         4: "Wave 4 — breadth and reach",
         5: "Wave O — Online mode (own API key), Private stays the default",
         6: "Wave N — new gaps from the market scan",
         7: "Wave I — innovation (no competitor has it)"}

# ---------------- Markdown
md = []
md.append("# Beat DEVONthink — final to-do list (Kalsmritikosh + mailin)\n")
md.append("Created 2026-10-15 (DEVONthink facts corrected the same day) from two inputs:\n"
          "1. the DEVONthink gap matrix (`~/Downloads/DEVONthink_vs_Kalsmritikosh_Gap_Matrix.xlsx`), and\n"
          "2. mailin's latest code (`~/Downloads/mailin`, branch `v3-architecture`, 17c925d) — what it already does that Kalsmritikosh lacks.\n")
md.append("## Strategy — two modes, toe to toe (owner decisions 2026-10-15)\n"
          "- **Private (default):** fully offline, proven: the main app has no network permission (W0-C).\n"
          "- **Online (opt-in):** the user's own Claude / OpenAI / Gemini key. Only the passages for one question are sent, shown before sending and logged. It runs in a separate helper process. No local model server, sync or MCP.\n"
          "- **Go toe to toe with DEVONthink on features**, not just a niche: close every table-stakes gap (Waves 1-4), porting from mailin where it already exists.\n"
          "- **Fight where the fight is fair and we can win: both apps offline.** DEVONthink's AI offline means a small local model the user installs and configures; ours is built in. The benchmark (W0-A) runs two rounds: both apps offline, then both on the same cloud model, so the ledger, not the model, decides.\n"
          "- **Proof before features.** W0 decides whether the moat is real: verified citations, honest refusals, contradictions, timelines.\n"
          "- **Make switching free:** import a DEVONthink database (W3-1) and work alongside it (W3-2).\n")
md.append("## Toe-to-toe scorecard — DEVONthink capability → our item\n"
          "| DEVONthink has | We answer with |\n|---|---|\n"
          "| AI chat (cloud or local model) | Private: built-in offline answers; Online: own-key Claude/OpenAI/Gemini (O-4) with send preview (O-5); same evidence gate both ways (O-6) |\n"
          "| MCP server for Claude/Codex | Not matched by design (it sends documents to an outside AI); Shortcuts (W1-6) for local automation |\n"
          "| Tags, smart groups, smart rules | W1-4, W3-5 |\n"
          "| Boolean/NEAR search | W1-1 (+ coverage badge W1-2, which DEVONthink lacks) |\n"
          "| Item links | W1-5 |\n"
          "| Annotation, Apple Pencil | W2-1 (overlay; sources never edited) |\n"
          "| OCR | W4-1 |\n"
          "| Versioning, revision-proof databases, audit log | Source versions (have) + signed receipts anyone can verify (W1-3) |\n"
          "| Stamps, Bates numbering | W2-5 (plus load files and legal holds, which DEVONthink lacks) |\n"
          "| Encrypted databases | W2-7 |\n"
          "| Email archiving | Have (PST/NSF/MSG FULL) + authenticity checks W2-2 + PST reader W2-3 |\n"
          "| Web clipping, Share/Services | W3-6 (saves locally; our app makes no connection) |\n"
          "| AppleScript/JXA | W1-6 Shortcuts first; AppleScript later |\n"
          "| iPhone/iPad app | W4-6 read-only companion from a signed snapshot file (no sync service) |\n"
          "| Sync between devices | Not in scope (owner, 2026-10-15); move the ledger on a disk (W1-8) |\n"
          "| Server/web access | Not matched by design |\n"
          "| 20 years of maturity, manual, community | W4-8 manual + 'Coming from DEVONthink' guide |\n"
          "| Price $99 / $199 / $499 | Planned $29-49 one-time (W4-9) |\n"
          "| — (none) | Contradictions, timelines, entity ledger, custody, persona workflows, forensic artifacts: our lead |\n")
md.append("## The innovation thesis\n"
          "Competitors answer *\u201cwhat do my files say?\u201d*. Kalsmritikosh answers four harder questions no one else does:\n"
          "1. **Is it true?** Corroboration meter (I-3), devil's-advocate check (I-4), measured accuracy on your own archive (I-7).\n"
          "2. **What is missing?** The gap map (I-2) and replies to emails we don't have (N-2).\n"
          "3. **What did we know, and when?** Time travel (I-5), statement drift (I-9), living dossiers (I-12), question monitors (N-11).\n"
          "4. **Can someone else check it?** Proof bundles anyone can verify (I-1), signed case slices (I-10), least-data online answers (I-8).\n")
md.append("## Build these first (the twelve that win the benchmark and the demo)\n"
          "| # | Item | Why first |\n|---|---|---|\n"
          "| 1 | W0-M module matrix | Everything after it lands behind a switch |\n"
          "| 2 | W0-C Private mode proven offline | Days of work; the privacy claim becomes checkable |\n"
          "| 3 | W0-A benchmark (round 1) | Tells us where we lose before we build more |\n"
          "| 4 | W0-B fix lost questions | The moat must be real |\n"
          "| 5 | I-3 corroboration meter | Small; visible on every answer |\n"
          "| 6 | I-4 devil's-advocate check | Turns honesty into a feature users see |\n"
          "| 7 | I-7 measured accuracy on your archive | No competitor can show this |\n"
          "| 8 | W1-1 + W1-2 search syntax + coverage badge | Table stakes, ported from mailin |\n"
          "| 9 | I-2 'what is missing' map | Builds on GapDetector, which already exists |\n"
          "| 10 | I-1 proof bundle | Upgrades receipts from consistency to authenticity |\n"
          "| 11 | O-1..O-6 Online mode (with I-8 least-data) | Round 2 of the benchmark |\n"
          "| 12 | N-1 + N-2 + N-3 chronology, communication, watchlists | Professional core no affordable private tool has |\n")
md.append("## Porting rules (from CLAUDE.md)\n"
          "- **Rewrite, don't link.** mailin's `ArchiveCore` has its own SQLite store and Swift 5.9. Port algorithms into Kalsmritikosh's `Database` actor, repositories and Swift 6 concurrency.\n"
          "- **Formats die at ingestion.** Email-specific signals (DKIM, ARC, S/MIME, timestamps, recovery) are computed in Ingestion and stored as ledger facts. Brain/Knowledge/Retrieval never branch on file type.\n"
          "- **Schema changes** only through new versioned migrations.\n"
          "- **No new third-party dependencies** without owner approval. SQLCipher (W2-7) is the only candidate.\n"
          "- **No network calls.** Spotlight and capture extensions stay local and are opt-in where they copy data.\n"
          "- **No deletion.** Near-duplicates and recovered items are linked and labelled, never merged away.\n")
md.append("Legend — **Source:** DEVONthink gap · Market gap (other competitors) · mailin port · both · Moat (DEVONthink can't) · Ship gate. **Effort:** S (days) · M (1–2 weeks) · L (3+ weeks). **Priority:** P0 blocks release … P3 nice to have.\n")
for w in [0, 5, 1, 7, 2, 3, 6, 4]:
    md.append(f"\n## {WAVES[w]}\n")
    md.append("| ID | Item | Source | Why | How (what to port) | Done means | Effort | Priority | Owner needed |")
    md.append("|---|---|---|---|---|---|---|---|---|")
    for t in [x for x in T if x[1] == w]:
        cells = [t[0], f"**{t[2]}**", t[3], t[4], t[5], t[6], t[7], t[8], t[9]]
        md.append("| " + " | ".join(c.replace("|", "\\|") for c in cells) + " |")
md.append("\n## Suggested order\n"
          "1. Follow **Build these first** above. In general: **W0-M first**, then **Wave 0 + Wave O.** O-1 (approve) and O-2 (helper target) unblock everything online. Then W0-C, then W0-A round 1. Build O-3..O-8, then run W0-A round 2. Loop W0-B until we win both rounds. Ship gates W0-1..3 run alongside.\n"
          "2. **Wave 1** next: W1-1, W1-2, W1-3, W1-5 and W1-4 first. They are small, ported and visible.\n"
          "3. Wave 2 P1 items: W2-1 annotations, W2-2 authenticity, W2-3 PST reader, W2-7 app lock.\n"
          "4. Wave 3 P1 items: W3-1 DEVONthink import, W3-3 visualisation, W3-4 alerts — plus Wave N P1 items N-1 chronology, N-2 communication analysis, N-3 watchlists, N-10 PII masking. These are the 'mile'.\n"
          "5. Then the remaining items by priority.\n")
md.append("## Decisions needed from the owner\n")
for t in T:
    if t[9].startswith("Yes") or t[9].startswith("Partly") or "check" in t[9].lower() or "review" in t[9].lower():
        md.append(f"- **{t[0]} {t[2]}:** {t[9]}")
md.append("\n## Market scan (2026-10-15)\n")
md.append("| Product | Kind | Price | Strengths | Weak spot vs us | Our answer |")
md.append("|---|---|---|---|---|---|")
for r in COMPETITORS:
    md.append("| " + " | ".join(r[:6]) + " |")
md.append("\nSources: " + ", ".join(f"[{r[0]}]({r[6]})" for r in COMPETITORS) + ".")
md.append("\n## Caveats\n"
          "- DEVONthink facts were checked on 2026-10-15 against devontechnologies.com (product page, release notes 4.0–4.4.1) and the MacStories DEVONthink 4 review.\n"
          "- Efforts are rough estimates.\n"
          "- mailin's measured results (e.g. the ARC suite, the 15 real PSTs) are mailin's own claims, recorded in its `FINAL_TODO_3_0_MAJOR.md`. Each one must be re-measured in Kalsmritikosh after porting.\n")
open(MD, "w").write("\n".join(md) + "\n")

# ---------------- Excel sheet (added to the gap-matrix workbook)
wb = openpyxl.load_workbook(XLSX)
if "Final TODO" in wb.sheetnames:
    del wb["Final TODO"]
ws = wb.create_sheet("Final TODO", 0)
hdr = ["ID", "Wave", "Item", "Source", "Why", "How (what to port)", "Done means", "Effort", "Priority", "Owner needed", "Status"]
ws.append(hdr)
for t in T:
    ws.append([t[0], WAVES[t[1]].split(" — ")[0], t[2], t[3], t[4], t[5], t[6], t[7], t[8], t[9], "To do"])
fill = PatternFill("solid", fgColor="1F3864"); thin = Side(style="thin", color="BFBFBF")
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = fill; c.alignment = Alignment(wrap_text=True, vertical="center")
for i, w in enumerate([7, 9, 34, 18, 42, 60, 44, 8, 9, 20, 10], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
n = len(T) + 1
for row in ws.iter_rows(min_row=2, max_row=n):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top"); c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
ws.freeze_panes = "D2"; ws.auto_filter.ref = f"A1:K{n}"
for p, col in {"P0": "F8CBAD", "P1": "FFE699", "P2": "DDEBF7", "P3": "E7E6E6"}.items():
    ws.conditional_formatting.add(f"I2:I{n}", FormulaRule(formula=[f'$I2="{p}"'], fill=PatternFill("solid", fgColor=col)))
ws.conditional_formatting.add(f"D2:D{n}", FormulaRule(formula=[f'$D2="{MOAT}"'], fill=PatternFill("solid", fgColor="C6EFCE")))
if "Competitors" in wb.sheetnames:
    del wb["Competitors"]
cs = wb.create_sheet("Competitors", 1)
cs.append(["Product", "Kind", "Price", "Strengths", "Weak spot vs us", "Our answer (plan IDs)", "Source"])
for r in COMPETITORS:
    cs.append(list(r))
for c in cs[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = fill; c.alignment = Alignment(wrap_text=True)
for i, w in enumerate([24, 26, 24, 60, 40, 28, 45], 1):
    cs.column_dimensions[get_column_letter(i)].width = w
for row in cs.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
cs.freeze_panes = "B2"
wb.active = 0
wb.save(XLSX)

from collections import Counter
print(MD); print(XLSX)
print(len(T), "items", Counter(t[8] for t in T), Counter(t[3] for t in T))
